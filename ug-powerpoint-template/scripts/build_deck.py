#!/usr/bin/env python3
"""Build a University of Ghana styled PowerPoint from a JSON specification.

Usage: python build_deck.py spec.json output.pptx
See SKILL.md for the schema and examples.
"""
import json
import sys
from pathlib import Path

from lxml import etree
from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
ASSETS = HERE.parent / "assets"
STANDARD_TEMPLATE = ASSETS / "UG_Powerpoint_template.pptx"
IMAGE_RICH_TEMPLATE = ASSETS / "UG Powerpoint Template - With Images_new.pptx"
TEMPLATE = STANDARD_TEMPLATE
TEMPLATE_OPTIONS = {
    "standard": STANDARD_TEMPLATE,
    "image_rich": IMAGE_RICH_TEMPLATE,
}
NAVY = RGBColor(0x28, 0x27, 0x81)
GOLD = RGBColor(0xD2, 0xAB, 0x67)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLIDE_W, SLIDE_H = 10693400, 7562850
BODY_TOP, BODY_BOTTOM = 2050000, 6600000
MARGIN_L, MARGIN_R = 639157, 10050000
MIN_BODY_PT = 18
MIN_SPLIT_PT = 16
BUILDERS = {}


def set_font(run, size, bold=True, italic=False, color=NAVY, face="Calibri"):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = face
    f.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for el in rpr.findall(qn("a:cs")):
        rpr.remove(el)
    cs = etree.SubElement(rpr, qn("a:cs"))
    cs.set("typeface", face)


def add_rect(slide, x, y, w, h, color, name):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    shp.name = name
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_textbox(slide, x, y, w, h, name):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb


def footer(slide):
    add_rect(slide, 8270514, 6711447, 2421890, 520700, WHITE, "Footer block")
    slide.shapes.add_picture(str(ASSETS / "logo_small.png"), Emu(8369007), Emu(6809791),
                             Emu(270236), Emu(324349)).name = "Footer crest"
    tb = add_textbox(slide, 8715147, 6899328, 1900000, 163829, "Footer text")
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = "UNIVERSITY OF GHANA"
    set_font(r, 10.5, face="Arial")


def style_bullet(p, level):
    ppr = p._p.get_or_add_pPr()
    ppr.set("marL", "469265" if level == 0 else "926465")
    ppr.set("indent", "-456565" if level == 0 else "-513715")
    if level:
        ppr.set("lvl", "1")
    char = "•" if level == 0 else "–"
    for tag in ("a:buFont", "a:buChar", "a:buNone"):
        for el in ppr.findall(qn(tag)):
            ppr.remove(el)
    bf = etree.SubElement(ppr, qn("a:buFont")); bf.set("typeface", "Arial")
    bc = etree.SubElement(ppr, qn("a:buChar")); bc.set("char", char)


def flatten(items):
    out = []
    for item in items:
        if isinstance(item, list):
            out.extend((1, value) for value in item)
        else:
            out.append((0, item))
    return out


def wrap_count(text, max_chars):
    """Conservative word-wrap estimate, accounting for long unbroken tokens."""
    words = text.split() or [""]
    lines, current = 1, 0
    for word in words:
        needed = len(word) + (1 if current else 0)
        if current and current + needed > max_chars:
            lines += 1
            current = len(word)
        else:
            current += needed
        if len(word) > max_chars:
            extra = (len(word) - 1) // max_chars
            lines += extra
            current = len(word) - extra * max_chars
    return lines


def fit_size(entries, width_emu, height_emu, start=32, floor=MIN_BODY_PT):
    width_pt = max(40, width_emu / 12700 - 37)
    for size in range(start, floor - 1, -2):
        total_pt = 0.0
        for level, text in entries:
            s = size if level == 0 else size - 4
            usable = max(20, width_pt - (36 if level else 0))
            chars = max(8, int(usable / (s * 0.48)))
            total_pt += wrap_count(text, chars) * s * 1.25 + s * 0.5
        if total_pt * 12700 <= height_emu:
            return size
    return floor


def fill_bullets(tb, entries, size):
    tf = tb.text_frame
    for i, (level, text) in enumerate(entries):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        style_bullet(p, level)
        p.space_before = Pt(size * 0.38)
        p.space_after = Pt(size * 0.12)
        r = p.add_run(); r.text = text
        set_font(r, size if level == 0 else max(12, size - 4))


def header(slide, title):
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = Emu(632128), Emu(640000), Emu(9500000), Emu(900000)
    tf = t.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.text = ""
    size = 44 if len(title) <= 30 else 36 if len(title) <= 44 else 30
    r = tf.paragraphs[0].add_run(); r.text = title
    set_font(r, size, color=WHITE, face="Arial")


def new_slide(prs, layout_name):
    layout = next((l for l in prs.slide_layouts if l.name == layout_name), None)
    if layout is None:
        raise ValueError(f"Required layout {layout_name!r} is missing from the bundled template.")
    return prs.slides.add_slide(layout)


def strip_unused_placeholders(slide, keep_title=True):
    for shp in list(slide.placeholders):
        if shp == slide.shapes.title and keep_title:
            continue
        shp._element.getparent().remove(shp._element)


def build_title(prs, spec):
    s = new_slide(prs, "Title Slide")
    s._element.set("showMasterSp", "0")
    add_rect(s, 0, 0, SLIDE_W, 4720226, NAVY, "Navy field")
    add_rect(s, 0, 4720226, SLIDE_W, 319405, GOLD, "Gold band")
    add_rect(s, 0, 5039444, SLIDE_W, SLIDE_H - 5039444, WHITE, "White field")
    title = next((shp for shp in s.placeholders if shp.placeholder_format.idx == 0), None)
    if title is None:
        raise ValueError("The template title slide has no title placeholder (index 0).")
    for shp in list(s.placeholders):
        if shp != title:
            shp._element.getparent().remove(shp._element)
    tree = s.shapes._spTree; tree.remove(title._element); tree.append(title._element)
    title.left, title.top, title.width, title.height = Emu(802005), Emu(1100000), Emu(9089390), Emu(2000000)
    tf = title.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.text = ""
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = spec["title"]
    n = len(spec["title"]); set_font(r, 44 if n <= 40 else 36 if n <= 80 else 30, color=WHITE, face="Arial")
    author = spec.get("author", "")
    if author:
        tb = add_textbox(s, 802005, 3400000, 9089390, 1100000, "Author")
        for i, line in enumerate(author.split("\n")):
            p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
            p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = line
            set_font(r, 28 if i == 0 else 20, color=WHITE, face="Arial")
    s.shapes.add_picture(str(ASSETS / "logo_large.png"), Emu(1347941), Emu(5642528),
                         Emu(1097474), Emu(1316949)).name = "UG crest"
    tb = add_textbox(s, 2792473, 5893355, 7200000, 699135, "University name")
    r = tb.text_frame.paragraphs[0].add_run(); r.text = "UNIVERSITY OF GHANA"
    set_font(r, 44, color=NAVY, face="Arial")
    return s


def build_bullets(prs, spec):
    s = new_slide(prs, "Title and Content"); strip_unused_placeholders(s); header(s, spec["title"])
    entries = flatten(spec["bullets"]); w, h = MARGIN_R - MARGIN_L, BODY_BOTTOM - BODY_TOP
    size = fit_size(entries, w, h)
    tb = add_textbox(s, MARGIN_L, BODY_TOP, w, h, "Body"); fill_bullets(tb, entries, size); footer(s)
    return s


def build_split(prs, spec):
    s = new_slide(prs, "Title and Content"); strip_unused_placeholders(s); header(s, spec["title"])
    lx, lw = 736974, 4300000; rx, rw = 5489770, MARGIN_R - 5489770
    h = BODY_BOTTOM - BODY_TOP; img = spec.get("image")
    if img:
        with Image.open(img) as im:
            iw, ih = im.size
        caption_h = 520000 if spec.get("left_text") else 0
        available_h = max(100000, h - caption_h - (120000 if caption_h else 0))
        scale = min(lw / iw, available_h / ih)
        w_, h_ = int(iw * scale), int(ih * scale)
        pic = s.shapes.add_picture(img, Emu(lx), Emu(BODY_TOP), Emu(w_), Emu(h_)); pic.name = "Figure"
        if spec.get("image_alt"):
            pic._element.nvPicPr.cNvPr.set("descr", spec["image_alt"])
        if spec.get("left_text"):
            tb = add_textbox(s, lx, BODY_TOP + h_ + 120000, lw, caption_h, "Caption")
            r = tb.text_frame.paragraphs[0].add_run(); r.text = spec["left_text"]
            set_font(r, 14, italic=True)
    elif spec.get("left_text"):
        tb = add_textbox(s, lx, BODY_TOP, lw, h, "Left text")
        size = 22 if len(spec["left_text"]) < 260 else 18
        for i, para in enumerate(spec["left_text"].split("\n")):
            p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
            p.space_after = Pt(8); r = p.add_run(); r.text = para; set_font(r, size, italic=True)
    y = BODY_TOP
    if spec.get("right_heading"):
        tb = add_textbox(s, rx, y, rw, 450000, "Sub heading")
        r = tb.text_frame.paragraphs[0].add_run(); r.text = spec["right_heading"]; set_font(r, 24)
        y += 560000
    if spec.get("right_bullets"):
        entries = flatten(spec["right_bullets"]); size = fit_size(entries, rw, BODY_BOTTOM - y, start=28, floor=MIN_SPLIT_PT)
        tb = add_textbox(s, rx, y, rw, BODY_BOTTOM - y, "Right bullets"); fill_bullets(tb, entries, size)
    footer(s); return s


BUILDERS.update({"title": build_title, "closing": build_title, "bullets": build_bullets, "split": build_split})


def require_text(value, label, allow_empty=False):
    if not isinstance(value, str) or (not allow_empty and not value.strip()):
        raise ValueError(f"{label} must be {'a string' if allow_empty else 'a non-empty string' }.")


def validate_spec(spec, spec_path):
    if not isinstance(spec, dict):
        raise ValueError("The JSON root must be an object containing a 'slides' array.")
    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ValueError("'slides' must be a non-empty array.")
    base = Path(spec_path).resolve().parent
    for i, sl in enumerate(slides, 1):
        label = f"Slide {i}"
        if not isinstance(sl, dict):
            raise ValueError(f"{label} must be a JSON object.")
        kind = sl.get("type", "bullets")
        if kind not in BUILDERS:
            raise ValueError(f"{label}: unknown type {kind!r}; choose one of {', '.join(BUILDERS)}.")
        require_text(sl.get("title"), f"{label} title")
        if kind in ("title", "closing"):
            if "author" in sl: require_text(sl["author"], f"{label} author", allow_empty=True)
        if kind == "bullets":
            bullets = sl.get("bullets")
            if not isinstance(bullets, list) or not bullets:
                raise ValueError(f"{label}: 'bullets' must be a non-empty array of strings or arrays of sub-bullets.")
            for j, item in enumerate(bullets, 1):
                vals = item if isinstance(item, list) else [item]
                if not vals:
                    raise ValueError(f"{label}: bullet item {j} is an empty sub-bullet array.")
                for value in vals: require_text(value, f"{label} bullet {j}")
        if kind == "split":
            for key in ("left_text", "right_heading", "image_alt"):
                if key in sl: require_text(sl[key], f"{label} {key}", allow_empty=True)
            bullets = sl.get("right_bullets", [])
            if not isinstance(bullets, list): raise ValueError(f"{label}: 'right_bullets' must be an array.")
            for j, item in enumerate(bullets, 1):
                vals = item if isinstance(item, list) else [item]
                if not vals: raise ValueError(f"{label}: right_bullets item {j} is empty.")
                for value in vals: require_text(value, f"{label} right bullet {j}")
            if not (sl.get("image") or sl.get("left_text") or sl.get("right_bullets") or sl.get("right_heading")):
                raise ValueError(f"{label}: split slide needs an image or some left/right content.")
            if sl.get("image"):
                image_path = Path(sl["image"])
                if not image_path.is_absolute(): image_path = base / image_path
                if not image_path.is_file(): raise ValueError(f"{label}: image file not found: {image_path}")
                try:
                    with Image.open(image_path) as im: im.verify()
                except (OSError, UnidentifiedImageError) as exc:
                    raise ValueError(f"{label}: image is unreadable or unsupported: {image_path} ({exc})") from exc
        if "notes" in sl: require_text(sl["notes"], f"{label} notes", allow_empty=True)
    template_choice = spec.get("template", "standard")
    if template_choice not in TEMPLATE_OPTIONS:
        raise ValueError("'template' must be either 'standard' or 'image_rich'.")
    selected_template = TEMPLATE_OPTIONS[template_choice]
    if not selected_template.is_file():
        raise ValueError(f"Selected template asset is missing: {selected_template}")
    for asset in (ASSETS / "UG_Powerpoint_template.pptx", ASSETS / "logo_small.png", ASSETS / "logo_large.png"):
        if not asset.is_file(): raise ValueError(f"Required bundled asset is missing: {asset}")


def main(spec_path, out_path):
    spec_file, output = Path(spec_path), Path(out_path)
    if not spec_file.is_file(): raise ValueError(f"Spec file not found: {spec_file}")
    try:
        spec = json.loads(spec_file.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"Invalid JSON at line {exc.lineno}, column {exc.colno}: {exc.msg}") from exc
    validate_spec(spec, spec_file)
    if output.suffix.lower() != ".pptx": raise ValueError("Output filename must end in .pptx.")
    output.parent.mkdir(parents=True, exist_ok=True)
    template_choice = spec.get("template", "standard")
    prs = Presentation(str(TEMPLATE_OPTIONS[template_choice]))
    for sld_id in list(prs.slides._sldIdLst):
        prs.part.drop_rel(sld_id.rId); prs.slides._sldIdLst.remove(sld_id)
    base = spec_file.resolve().parent
    for i, sl in enumerate(spec["slides"], 1):
        kind = sl.get("type", "bullets")
        if sl.get("image") and not Path(sl["image"]).is_absolute():
            sl["image"] = str(base / sl["image"])
        try:
            slide = BUILDERS[kind](prs, sl)
            if sl.get("notes"): slide.notes_slide.notes_text_frame.text = sl["notes"]
        except Exception as exc:
            raise RuntimeError(f"Failed to build slide {i} ({kind}): {exc}") from exc
    prs.save(str(output))
    check = Presentation(str(output))
    if len(check.slides) != len(spec["slides"]):
        raise RuntimeError(f"Saved deck has {len(check.slides)} slides; expected {len(spec['slides'])}.")
    print(f"Wrote {output} ({len(check.slides)} slides). Visual inspection is still required.")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    try:
        main(sys.argv[1], sys.argv[2])
    except (ValueError, RuntimeError, OSError) as exc:
        raise SystemExit(f"ERROR: {exc}")
