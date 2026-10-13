import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
BUILDER = ROOT / 'scripts' / 'build_deck.py'

class BuilderTests(unittest.TestCase):
    def run_builder(self, folder, spec, name='out.pptx'):
        spec_path = folder / 'spec.json'
        spec_path.write_text(json.dumps(spec), encoding='utf-8')
        out = folder / name
        result = subprocess.run([sys.executable, str(BUILDER), str(spec_path), str(out)], capture_output=True, text=True)
        return result, out

    def test_all_supported_slide_types(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            spec = {'slides': [
                {'type':'title','title':'Test presentation','author':'Presenter\nUniversity of Ghana'},
                {'type':'bullets','title':'Main findings','bullets':['A grounded point',['A sub point','Another sub point'],'Final point'],'notes':'Presenter notes'},
                {'type':'split','title':'Figure and findings','left_text':'A short caption','image':str(ROOT/'assets'/'logo_large.png'),'image_alt':'University crest','right_heading':'Findings','right_bullets':['Finding one','Finding two']},
                {'type':'closing','title':'Thank you','author':'Questions'}
            ]}
            result, out = self.run_builder(folder, spec)
            self.assertEqual(result.returncode, 0, result.stderr)
            prs = Presentation(str(out))
            self.assertEqual(len(prs.slides), 4)
            self.assertTrue(any('Main findings' in sh.text for sh in prs.slides[1].shapes if sh.has_text_frame))

    def test_missing_bullets_fails_clearly(self):
        with tempfile.TemporaryDirectory() as tmp:
            result, _ = self.run_builder(Path(tmp), {'slides':[{'type':'bullets','title':'Missing bullets'}]})
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("'bullets' must be a non-empty array", result.stderr)

    def test_unknown_type_fails_clearly(self):
        with tempfile.TemporaryDirectory() as tmp:
            result, _ = self.run_builder(Path(tmp), {'slides':[{'type':'chart','title':'Unsupported'}]})
            self.assertNotEqual(result.returncode, 0)
            self.assertIn('unknown type', result.stderr)

    def test_missing_image_fails_clearly(self):
        with tempfile.TemporaryDirectory() as tmp:
            result, _ = self.run_builder(Path(tmp), {'slides':[{'type':'split','title':'Bad image','image':'missing.png'}]})
            self.assertNotEqual(result.returncode, 0)
            self.assertIn('image file not found', result.stderr)

    def test_invalid_json_fails_clearly(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp); path = folder/'spec.json'; path.write_text('{bad json', encoding='utf-8')
            result = subprocess.run([sys.executable, str(BUILDER), str(path), str(folder/'out.pptx')], capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn('Invalid JSON', result.stderr)

if __name__ == '__main__':
    unittest.main(verbosity=2)
